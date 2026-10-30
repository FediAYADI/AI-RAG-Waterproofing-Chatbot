"""
Document Processor - Handles text extraction from various file formats
"""

import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from typing import List, Tuple
import pypdf
from docx import Document
import pandas as pd
from pathlib import Path


class DocumentProcessor:
    """Process documents and extract text from various formats"""

    def __init__(self):
        self.supported_extensions = {
            '.pdf': self._extract_pdf,
            '.txt': self._extract_txt,
            '.docx': self._extract_docx,
            '.md': self._extract_txt,
            '.csv': self._extract_csv,
            '.xls': self._extract_excel,
            '.xlsx': self._extract_excel,
        }

    def extract_text(self, file_path: str) -> str:
        """
        Extract text from a document based on its file extension
        
        Args:
            file_path: Path to the document file
            
        Returns:
            Extracted text as a string
            
        Raises:
            ValueError: If file format is not supported
            Exception: If extraction fails
        """
        file_ext = Path(file_path).suffix.lower()
        
        if file_ext not in self.supported_extensions:
            raise ValueError(f"Unsupported file format: {file_ext}")
        
        extractor = self.supported_extensions[file_ext]
        return extractor(file_path)

    def _extract_pdf(self, file_path: str) -> str:
        """Extract text from PDF files"""
        text = ""
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = pypdf.PdfReader(file)
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
        except Exception as e:
            raise Exception(f"Failed to extract text from PDF: {str(e)}")
        return text

    def _extract_txt(self, file_path: str) -> str:
        """Extract text from plain text and markdown files"""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                return file.read()
        except UnicodeDecodeError:
            # Try with different encoding
            try:
                with open(file_path, 'r', encoding='latin-1') as file:
                    return file.read()
            except Exception as e:
                raise Exception(f"Failed to read text file: {str(e)}")
        except Exception as e:
            raise Exception(f"Failed to read text file: {str(e)}")

    def _extract_docx(self, file_path: str) -> str:
        """Extract text from Word documents (.docx)"""
        try:
            doc = Document(file_path)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
                text += "\n"
            return text
        except Exception as e:
            raise Exception(f"Failed to extract text from DOCX: {str(e)}")

    def _extract_csv(self, file_path: str) -> str:
        """
        Extract text from CSV files
        Converts CSV data to readable text format
        """
        try:
            df = pd.read_csv(file_path)
            return self._dataframe_to_text(df)
        except Exception as e:
            raise Exception(f"Failed to extract text from CSV: {str(e)}")

    def _extract_excel(self, file_path: str) -> str:
        """
        Extract text from Excel files (.xls, .xlsx)
        Reads all sheets and converts to text format
        """
        try:
            # Read all sheets
            excel_file = pd.ExcelFile(file_path)
            text = ""
            
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet_name)
                text += f"\n=== Sheet: {sheet_name} ===\n"
                text += self._dataframe_to_text(df)
                text += "\n"
            
            return text
        except Exception as e:
            raise Exception(f"Failed to extract text from Excel: {str(e)}")

    def _dataframe_to_text(self, df: pd.DataFrame) -> str:
        """
        Convert pandas DataFrame to readable text format
        
        Args:
            df: Pandas DataFrame
            
        Returns:
            Formatted text representation of the data
        """
        text = ""
        
        # Add column headers
        if not df.empty:
            text += "Columns: " + ", ".join(str(col) for col in df.columns) + "\n\n"
            
            # Convert each row to text
            for index, row in df.iterrows():
                row_text = " | ".join(f"{col}: {val}" for col, val in row.items())
                text += row_text + "\n"
        
        return text

    def chunk_text(self, text: str, chunk_size: int = 1000, overlap: int = 200) -> List[str]:
        """
        Split text into chunks for embedding and vector storage
        
        Args:
            text: Input text to chunk
            chunk_size: Maximum size of each chunk
            overlap: Number of characters to overlap between chunks
            
        Returns:
            List of text chunks
        """
        chunks = []
        start = 0
        text_length = len(text)
        
        while start < text_length:
            end = start + chunk_size
            chunk = text[start:end]
            
            # Try to break at sentence boundaries
            if end < text_length:
                last_period = chunk.rfind('.')
                last_newline = chunk.rfind('\n')
                last_break = max(last_period, last_newline)
                
                if last_break > chunk_size * 0.5:  # Only break if it's not too early
                    chunk = chunk[:last_break + 1]
                    end = start + last_break + 1
            
            chunks.append(chunk.strip())
            start = end - overlap
        
        return chunks

    def process_document(self, file_path: str, chunk_size: int = 1000, overlap: int = 200) -> Tuple[str, List[str]]:
        """
        Complete document processing pipeline: extract text and chunk it
        
        Args:
            file_path: Path to the document
            chunk_size: Maximum chunk size
            overlap: Chunk overlap
            
        Returns:
            Tuple of (full_text, chunks)
        """
        text = self.extract_text(file_path)
        chunks = self.chunk_text(text, chunk_size, overlap)
        return text, chunks


# Singleton instance
processor = DocumentProcessor()
